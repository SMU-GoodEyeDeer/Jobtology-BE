import time
from datetime import UTC, datetime, timedelta
from itertools import pairwise
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import anyio
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from fastapi.testclient import TestClient
from sqlalchemy.exc import SQLAlchemyError
from test_product_roles_approval import _inputs
from test_recompute_worker import (
    RecordingFinalizer,
    StaticContextReader,
    StaticSnapshotReader,
    StaticWorkClaimer,
    _context,
    _snapshot,
    _work_item,
)

from jobtology_be.api.composition import database_lifespan
from jobtology_be.api.readiness import RuntimeReadiness
from jobtology_be.api.router import create_api_shell, router
from jobtology_be.infrastructure.persistence.contracts import PersistenceConflictError
from jobtology_be.main import create_app
from jobtology_be.planning.calendar import flexible_weekly_slots
from jobtology_be.planning.solver_models import PlanningSlot
from jobtology_be.product_roles.holder import ProductRoleHolder
from jobtology_be.product_roles.worker import InProcessRecomputeLoop, WorkerHealth
from jobtology_be.resume.worker import visible_docx
from jobtology_be.settings import Settings
from jobtology_be.workers.recompute import LeasedRecomputeWorker, RecomputeWorker


@pytest.mark.anyio
async def test_failure_finalization_lease_loss_does_not_stop_later_items():
    class FailureFinalizer:
        calls = 0

        async def fail_recompute(self, work_item, error_code):
            self.calls += 1
            if self.calls == 1:
                raise PersistenceConflictError(resource="outbox lease")

    failures = FailureFinalizer()
    worker = LeasedRecomputeWorker(
        StaticWorkClaimer((_work_item(), _work_item())),
        RecomputeWorker(StaticContextReader(_context(profile_version=4)),
                        StaticSnapshotReader(_snapshot()), RecordingFinalizer()),
        failures,
    )
    assert await worker.process_once(limit=2) == ()
    assert failures.calls == 2


@pytest.mark.anyio
async def test_worker_recovers_from_unexpected_cycle_failure_and_reports_progress(monkeypatch):
    failures_seen = []
    recovered = anyio.Event()

    async def process_once(self):
        if not failures_seen:
            failures_seen.append(True)
            raise ValueError("unexpected work-item error")
        assert self.health.last_cycle_failed
        recovered.set()

    monkeypatch.setattr(InProcessRecomputeLoop, "process_once", process_once)
    monkeypatch.setattr("jobtology_be.product_roles.worker.RECOMPUTE_CYCLE_DELAY_SECONDS", 0)
    worker = InProcessRecomputeLoop(None, None, None, None)
    async with anyio.create_task_group() as tasks:
        tasks.start_soon(worker.run)
        await recovered.wait()
        assert worker.health.ready
        tasks.cancel_scope.cancel()
    assert not worker.health.ready


@pytest.mark.anyio
async def test_deferred_roadmap_follow_up_preserves_other_accounts_and_fails_cycle():
    class Worker:
        async def process_once(self, **kwargs):
            return ()

    class Session:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return None

        async def execute(self, query):
            return self

        def all(self):
            return [SimpleNamespace(id=number, occupation_id=None) for number in (1, 2)]

    class Roadmaps:
        def __init__(self):
            self.calls = []

        async def reconcile_automatic(self, request_id, title):
            self.calls.append(request_id)
            if request_id == 1:
                raise SQLAlchemyError("transient follow-up error")

    roadmaps = Roadmaps()
    loop = InProcessRecomputeLoop(Worker(), SimpleNamespace(sessions=Session), roadmaps, None)
    with pytest.raises(SQLAlchemyError):
        await loop.process_once()
    assert roadmaps.calls == [1, 2]


@pytest.mark.anyio
async def test_readiness_requires_a_recent_cycle_and_a_running_worker():
    class Worker:
        health = WorkerHealth(running=True, last_cycle=time.monotonic())

    worker = Worker()
    state = RuntimeReadiness(None, None, worker)
    assert await state.ready()
    worker.health.last_cycle -= worker.health.stale_after_seconds + 1
    assert not await state.ready()
    worker.health.last_cycle = time.monotonic()
    worker.health.running = False
    assert not await state.ready()


@pytest.mark.anyio
async def test_worker_tolerates_isolated_failures_and_recovers_after_failure_threshold(monkeypatch, caplog):
    observations = []
    recovered = anyio.Event()
    calls = 0

    async def process_once(self):
        nonlocal calls
        if calls:
            observations.append((self.health.last_cycle, self.health.consecutive_failures,
                                 self.health.ready))
        calls += 1
        if calls <= self.health.failure_threshold:
            raise ValueError('work-item failure')
        if calls > self.health.failure_threshold + 1:
            recovered.set()
            await anyio.sleep_forever()

    monkeypatch.setattr(InProcessRecomputeLoop, 'process_once', process_once)
    monkeypatch.setattr('jobtology_be.product_roles.worker.RECOMPUTE_CYCLE_DELAY_SECONDS', 0)
    worker = InProcessRecomputeLoop(None, None, None, None)
    assert not worker.health.ready
    async with anyio.create_task_group() as tasks:
        tasks.start_soon(worker.run)
        await recovered.wait()
        assert [(count, ready) for _, count, ready in observations] == [
            (1, True), (2, True), (3, False), (0, True),
        ]
        assert all(stamp is not None for stamp, _, _ in observations)
        assert [stamp for stamp, _, _ in observations] == sorted(stamp for stamp, _, _ in observations)
        assert worker.health.last_success is not None
        tasks.cancel_scope.cancel()
    assert not worker.health.ready
    assert caplog.text.count('product role recompute cycle deferred after error') == 3


@pytest.mark.parametrize("retryable", [True, False])
def test_catalog_startup_recovery_keeps_unapproved_catalog_closed(monkeypatch, retryable):
    class Holder:
        available = False
        retryable_load_failure = retryable
        calls = 0

        async def load(self):
            self.calls += 1
            self.available = retryable and self.calls > 1
            return self.available

    holder = Holder()
    monkeypatch.setattr("jobtology_be.api.composition.CATALOG_RETRY_DELAY_SECONDS", 0.01)
    app = create_api_shell(database_lifespan(None, role_holder=holder))
    app.include_router(router, prefix="/api/v1")
    with TestClient(app) as client:
        assert client.get("/api/v1/health/live").status_code == 200
        deadline = time.monotonic() + 1
        response = client.get("/api/v1/health/ready")
        while retryable and response.status_code == 503 and time.monotonic() < deadline:
            time.sleep(0.01)
            response = client.get("/api/v1/health/ready")
        assert response.status_code == (200 if retryable else 503)
        if not retryable:
            time.sleep(0.03)
            assert holder.calls == 1


@pytest.mark.anyio
@pytest.mark.parametrize('failure', [ConnectionRefusedError, OSError])
async def test_catalog_connection_failure_is_retryable(failure):
    class Engine:
        def connect(self):
            raise failure("offline")

    policy = Path(__file__).parents[1] / "config/product_roles/policy.v1.json"
    holder = ProductRoleHolder(Engine(), policy)
    assert not await holder.load()
    assert holder.retryable_load_failure
    assert not holder.available


@pytest.mark.anyio
async def test_catalog_startup_connection_attempt_has_a_deadline(monkeypatch):
    class Connection:
        async def __aenter__(self):
            await anyio.sleep_forever()

        async def __aexit__(self, *args):
            return None

    class Engine:
        def connect(self):
            return Connection()

    monkeypatch.setattr("jobtology_be.product_roles.holder.ROLE_LOAD_TIMEOUT_SECONDS", 0.01)
    holder = ProductRoleHolder(Engine(), Path(__file__).parents[1]
                               / "config/product_roles/policy.v1.json")
    assert not await holder.load()
    assert holder.retryable_load_failure
    assert not holder.available


@pytest.mark.anyio
async def test_catalog_digest_mismatch_is_unavailable_and_not_automatically_retried(tmp_path):
    class Connection:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return None

        def begin(self):
            return self

        async def execute(self, *args):
            return self

        def scalar_one(self):
            return _inputs().model_dump_json()

    class Engine:
        def connect(self):
            return Connection()

    approval = tmp_path / "approval.json"
    approval.write_text('{"sha256":"' + '0' * 64 + '","approved_by":"test",'
                        '"reviewed_at":"2026-10-05T09:00:00+00:00"}')
    holder = ProductRoleHolder(Engine(), Path(__file__).parents[1]
                               / "config/product_roles/policy.v1.json", approval)
    assert not await holder.load()
    assert holder.draft is not None
    assert not holder.available
    assert not holder.retryable_load_failure


@pytest.mark.anyio
async def test_readiness_reports_database_outage():
    class Database:
        def sessions(self):
            raise ConnectionRefusedError("offline")

    assert not await RuntimeReadiness(Database(), None, None).ready()


@pytest.mark.parametrize("required_setting", ["product_roles_enabled", "inprocess_worker_enabled"])
def test_readiness_rejects_required_but_unconfigured_runtime_component(required_setting):
    app = create_app(Settings(_env_file=None, **{required_setting: True}))
    with TestClient(app) as client:
        assert client.get("/api/v1/health/live").status_code == 200
        assert client.get("/api/v1/health/ready").status_code == 503


@pytest.mark.parametrize(("month", "day", "expected"), [(3, 8, 3), (11, 1, 5)])
def test_dst_calendar_has_real_hour_slots_after_timestamp_serialization(month, day, expected):
    timezone = ZoneInfo("America/New_York")
    start = datetime(2026, month, day, tzinfo=timezone)
    end = start.replace(hour=4)
    slots = flexible_weekly_slots(start, end, 20)
    assert len(slots) == expected
    replayed = tuple(PlanningSlot(datetime.fromisoformat(slot.starts_at.isoformat()),
                                  datetime.fromisoformat(slot.ends_at.isoformat()),
                                  slot.capacity_week_key) for slot in slots)
    assert all(slot.ends_at - slot.starts_at == timedelta(hours=1) for slot in replayed)
    assert all(left.ends_at == right.starts_at for left, right in pairwise(replayed))
    assert replayed[-1].ends_at == end.astimezone(UTC)


def test_calendar_capacity_week_is_local_even_though_slots_are_utc():
    timezone = ZoneInfo("Pacific/Kiritimati")
    start = datetime(2026, 10, 18, 23, tzinfo=timezone)
    slots = flexible_weekly_slots(start, start + timedelta(hours=3), 1)
    assert [slot.capacity_week_key for slot in slots] == ["2026-W42", "2026-W43"]


def test_docx_hidden_style_false_overrides_and_direct_run_visibility(tmp_path):
    document = Document()
    hidden = document.styles.add_style("HiddenBase", WD_STYLE_TYPE.CHARACTER)
    hidden.font.hidden = True
    visible = document.styles.add_style("VisibleChild", WD_STYLE_TYPE.CHARACTER)
    visible.base_style = hidden
    visible.font.hidden = False
    paragraph = document.add_paragraph()
    paragraph.add_run("HIDDEN_STYLE", style=hidden)
    paragraph.add_run("VISIBLE_CHILD", style=visible)
    paragraph.add_run("VISIBLE_DIRECT", style=hidden).font.hidden = False
    hidden_paragraph = document.styles.add_style("HiddenParagraph", WD_STYLE_TYPE.PARAGRAPH)
    hidden_paragraph.font.hidden = True
    paragraph = document.add_paragraph(style=hidden_paragraph)
    paragraph.add_run("HIDDEN_PARAGRAPH")
    paragraph.add_run("VISIBLE_PARAGRAPH_OVERRIDE").font.hidden = False
    source = tmp_path / "input.docx"
    document.save(source)
    output, removed = visible_docx(source)
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert removed
    assert "HIDDEN_STYLE" not in text and "HIDDEN_PARAGRAPH" not in text
    assert "VISIBLE_CHILD" in text and "VISIBLE_DIRECT" in text
    assert "VISIBLE_PARAGRAPH_OVERRIDE" in text
