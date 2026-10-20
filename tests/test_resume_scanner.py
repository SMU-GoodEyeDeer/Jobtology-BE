import asyncio
import fcntl
import hashlib
import io
import os
import time
import zipfile
from uuid import UUID, uuid4

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from fastapi.testclient import TestClient
from pydantic import ValidationError
from pypdf import PdfWriter
from test_auth_session_api import FRONTEND_ORIGIN, RecordingSessionStore, _session_cookie
from test_capabilities_api import StaticIdentityProvider
from test_specs import NCS, catalog

from jobtology_be.api.dependencies import ApiDependencies
from jobtology_be.application.specs import DocumentEvidence, SpecDetails, capability_inputs
from jobtology_be.infrastructure.persistence.rate_limit import RateLimitDecision
from jobtology_be.main import create_app
from jobtology_be.modules.auth.session import issue_session_credentials
from jobtology_be.resume.matching import literal_suggestions, validated_ai_picks
from jobtology_be.resume.models import MAX_BYTES, PARSER_REVISION, ResumeBlock
from jobtology_be.resume.service import ResumeError, ResumeService
from jobtology_be.resume.worker import DocumentRejected, preflight
from jobtology_be.settings import Settings

TEXT = '😀 Python과 FastAPI로 로그인 API 프로젝트를 직접 구현하고 테스트했습니다.'
USER = UUID(int=42)


def block(text=TEXT):
    return ResumeBlock(block_id='a' * 24, text=text, text_hash=hashlib.sha256(text.encode()).hexdigest())


def docx_bytes():
    stream = io.BytesIO()
    doc = Document()
    doc.add_paragraph(TEXT)
    doc.add_paragraph().add_run('HIDDEN_TOKEN').font.hidden = True
    style = doc.styles.add_style('HiddenCharacter', WD_STYLE_TYPE.CHARACTER)
    style.font.hidden = True
    paragraph = doc.add_paragraph('Visible alongside hidden style')
    paragraph.add_run('HIDDEN_STYLE_TOKEN', style='HiddenCharacter')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '테이블에서 Python 데이터 분석을 직접 수행했습니다.'
    doc.add_paragraph('TOEIC 850 / 990')
    doc.save(stream)
    return stream.getvalue()


class Limiter:
    def __init__(self, allowed=True):
        self.allowed = allowed
        self.calls = []

    async def consume(self, user_id):
        self.calls.append(user_id)
        return RateLimitDecision(allowed=self.allowed, retry_after_seconds=60)


def application(tmp_path, limiter=None, session=None, llm=None):
    service = ResumeService(catalog(), root=tmp_path / 'uploads', formats=('docx', 'hwpx', 'pdf'), llm=llm)
    return create_app(
        Settings(_env_file=None, environment='test', cors_origins=[FRONTEND_ORIGIN]),
        dependencies=ApiDependencies(
            identity_provider=None if session else StaticIdentityProvider(),
            session_store=session, resume_service=service, chat_rate_limiter=limiter or Limiter(),
        ),
    ), service


def test_real_authenticated_scan_includes_table_text_and_exact_evidence_without_ai_or_writes(tmp_path):
    limiter = Limiter()
    app, service = application(tmp_path, limiter)
    content = docx_bytes()
    with TestClient(app) as client:
        status = client.get('/api/v1/me/resume/status')
        response = client.post('/api/v1/me/resume/scan?format=docx', content=content)
    assert status.json()['retention'] == 'REQUEST_ONLY'
    assert status.json()['ai_available'] is False
    assert response.status_code == 200, response.text
    result = response.json()
    assert response.headers['cache-control'] == 'no-store'
    assert result['document_sha256'] == hashlib.sha256(content).hexdigest()
    text = '\n'.join(item['text'] for item in result['blocks'])
    assert TEXT in text
    assert 'HIDDEN_TOKEN' not in text and 'HIDDEN_STYLE_TOKEN' not in text
    assert 'Visible alongside hidden style' in text
    assert text.count('테이블에서') == 1
    assert all(item['page'] is None for item in result['blocks'])
    assert 'HIDDEN_TEXT_REMOVED' in result['warnings']
    assert result['suggestions'][0]['competencies'][0]['entity_id'] == NCS
    assert result['suggestions'][0]['quote'] == TEXT
    assert result['suggestions'][0]['end'] == len(TEXT)
    assert any(item['score'] == 850 and item['score_max'] == 990 for item in result['suggestions'])
    assert not list(service.root.glob('scan-*'))
    assert len(limiter.calls) == 1


def test_authentication_and_csrf_reject_before_parsing(tmp_path):
    credentials = issue_session_credentials()
    app, service = application(tmp_path, session=RecordingSessionStore(credentials))
    with TestClient(app) as client:
        unauth = client.post('/api/v1/me/resume/scan?format=docx', content=docx_bytes())
        rejected = client.post('/api/v1/me/resume/scan?format=docx', content=docx_bytes(),
                               cookies=_session_cookie(credentials), headers={'Origin': FRONTEND_ORIGIN})
        accepted = client.post('/api/v1/me/resume/scan?format=docx', content=docx_bytes(),
                               cookies=_session_cookie(credentials), headers={
                                   'Origin': FRONTEND_ORIGIN, 'X-CSRF-Token': credentials.csrf_token})
    assert unauth.status_code == 401
    assert rejected.status_code == 403
    assert accepted.status_code == 200
    assert not list(service.root.glob('scan-*'))


@pytest.mark.parametrize(('content', 'expected'), [(b'not a document', 'FORMAT_MISMATCH'), (b'', 'EMPTY_DOCUMENT')])
def test_bad_documents_are_safe_errors_and_files_are_deleted(tmp_path, content, expected, caplog):
    app, service = application(tmp_path)
    with TestClient(app) as client:
        response = client.post('/api/v1/me/resume/scan?format=docx', content=content)
    assert response.status_code == 422
    assert response.json()['error']['code'] == expected
    request_id = response.json()['error']['request_id']
    assert response.headers['x-request-id'] == request_id
    assert f'request_id={request_id}' in caplog.text
    assert f'code={expected}' in caplog.text
    assert str(tmp_path) not in caplog.text and 'not a document' not in caplog.text
    assert not list(service.root.glob('scan-*'))


def test_size_limit_and_account_quota(tmp_path):
    app, service = application(tmp_path)
    with TestClient(app) as client:
        response = client.post('/api/v1/me/resume/scan?format=docx', content=b'a', headers={'content-length': str(MAX_BYTES + 1)})
    assert response.status_code == 413
    assert not list(service.root.glob('scan-*'))
    app, _ = application(tmp_path, Limiter(False))
    with TestClient(app) as client:
        response = client.post('/api/v1/me/resume/scan?format=docx', content=docx_bytes())
    assert response.status_code == 429
    assert response.headers['retry-after'] == '60'


def test_streamed_size_limit_cleans_partial_upload(tmp_path):
    service = ResumeService(catalog(), root=tmp_path / 'uploads', formats=('docx',))
    async def stream():
        yield b'a' * (MAX_BYTES // 2)
        yield b'b' * MAX_BYTES
    with pytest.raises(ResumeError, match='FILE_TOO_LARGE'):
        asyncio.run(service.scan(USER, 'docx', stream()))
    assert not list(service.root.glob('scan-*'))


def test_slots_are_shared_between_service_instances_and_bounded_per_account(tmp_path):
    first = ResumeService(catalog(), root=tmp_path)
    second = ResumeService(catalog(), root=tmp_path)
    with first.slot(USER):
        with pytest.raises(ResumeError, match='SCAN_ALREADY_RUNNING'), second.slot(USER):
            pytest.fail('same account acquired twice')
        with second.slot(uuid4()), pytest.raises(ResumeError, match='SCANNER_BUSY'), first.slot(uuid4()):
            pytest.fail('third slot')
    with second.slot(USER):
        pass
    assert len(list(tmp_path.glob('*.lock'))) == 3


def test_crash_cleanup_skips_live_locks_and_deletes_stale_private_files(tmp_path):
    service = ResumeService(catalog(), root=tmp_path)
    stale = tmp_path / 'scan-stale'; stale.mkdir(); (stale / 'private.docx').write_bytes(b'private')
    live = tmp_path / 'scan-live'; live.mkdir()
    with (live / '.lock').open('a+') as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        old = time.time() - 1000
        os.utime(stale, (old, old)); os.utime(live, (old, old))
        service.cleanup()
        assert not stale.exists() and live.exists()
    service.cleanup()
    assert not live.exists()


def test_zip_limits_and_entities_are_rejected(tmp_path):
    for name, data, expected in [('word/document.xml', b'<!DOCTYPE x [<!ENTITY x "secret">]><x/>', 'UNSAFE_XML'),
                                 ('../word/document.xml', b'<x/>', 'ARCHIVE_LIMIT')]:
        path = tmp_path / 'unsafe.docx'
        with zipfile.ZipFile(path, 'w') as z:
            z.writestr('word/document.xml', data if name == 'word/document.xml' else b'<x/>')
            if name != 'word/document.xml':
                z.writestr(name, data)
        with pytest.raises(DocumentRejected, match=expected):
            preflight(path)


def test_pdf_page_limit_encryption_and_image_only_ocr_notice(tmp_path):
    path = tmp_path / 'input.pdf'
    writer = PdfWriter()
    for _ in range(21):
        writer.add_blank_page(width=100, height=100)
    writer.write(path)
    with pytest.raises(DocumentRejected, match='PAGE_LIMIT'):
        preflight(path)
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100); writer.encrypt('password'); writer.write(path)
    with pytest.raises(DocumentRejected, match='ENCRYPTED_DOCUMENT'):
        preflight(path)


def test_keyword_boundaries_do_not_turn_javascript_into_java():
    item = catalog().items()[0].model_copy(update={'aliases': ('Java',), 'label': 'Java', 'original_label': 'Java'})
    assert not literal_suggestions((block('JavaScript로 API를 구현했습니다.'),), (item,))[0]
    assert literal_suggestions((block('Java로 API를 구현했습니다.'),), (item,))[0]


def test_ai_rejects_invented_quotes_ids_offsets_and_handles_repeated_unicode_spans():
    text = '😀 Python Python'
    good = {'title': 'Python', 'category': 'skill', 'competency_ids': [NCS], 'quote': 'Python', 'start': 9, 'end': 15}
    values = [good, {**good, 'quote': 'invent'}, {**good, 'competency_ids': ['invented']},
              {**good, 'start': 2, 'end': 8, 'quote': 'Python'}, {**good, 'start': True}]
    result = validated_ai_picks(values, text, catalog().items())
    assert len(result.suggestions) == 2
    assert result.suggestions[0].start == 9
    assert result.rejected_count == 3


def test_ai_is_explicit_and_bounded_to_the_reviewed_passage(tmp_path):
    class Llm:
        def __init__(self):
            self.calls = []
        async def complete_json(self, **kwargs):
            self.calls.append(kwargs)
            return {'candidates': [{'title': 'Python 구현', 'category': 'project', 'competency_ids': [NCS],
                                    'quote': 'Python', 'start': 2, 'end': 8}]}
    llm = Llm()
    app, _ = application(tmp_path, llm=llm)
    with TestClient(app) as client:
        assert client.post('/api/v1/me/resume/scan?format=docx', content=docx_bytes()).status_code == 200
        assert not llm.calls
        response = client.post('/api/v1/me/resume/interpret', json={'text': TEXT})
        assert response.status_code == 200
        assert len(response.json()['suggestions']) == 1
    assert len(llm.calls) == 1
    assert llm.calls[0]['messages'][0].text == TEXT


def test_document_provenance_is_bounded_self_report_and_preserves_exact_quote():
    source = DocumentEvidence(document_sha256='a' * 64, parser_revision=PARSER_REVISION,
                              block_id='b' * 24, text_hash='c' * 64,
                              quote=TEXT, start=0, end=len(TEXT), method='KEYWORD')
    details = SpecDetails(source='DOCUMENT_IMPORT', document_evidence=source, competency_ids=(NCS,), evidence=TEXT)
    assert capability_inputs(raw_text='Python', entity_id=None, category='project', proficiency='GUIDED',
                             details=details.model_dump())[0].credit == .5
    with pytest.raises(ValidationError):
        SpecDetails(source='DOCUMENT_IMPORT')
    with pytest.raises(ValidationError):
        DocumentEvidence(**{**source.model_dump(), 'end': len(TEXT) + 1})


def test_parser_timeout_kills_child_and_erases_source(tmp_path, monkeypatch):
    import sys

    import jobtology_be.resume.service as module
    original = asyncio.create_subprocess_exec
    children = []
    async def slow_child(*args, **kwargs):
        process = await original(sys.executable, '-c', 'import time; time.sleep(30)', **kwargs)
        children.append(process)
        return process
    monkeypatch.setattr(module, 'PARSE_SECONDS', .1)
    monkeypatch.setattr(module.asyncio, 'create_subprocess_exec', slow_child)
    service = ResumeService(catalog(), root=tmp_path, formats=('docx',))
    async def stream():
        yield docx_bytes()
    with pytest.raises(ResumeError, match='PARSE_TIMEOUT'):
        asyncio.run(service.scan(USER, 'docx', stream()))
    assert children[0].returncode < 0
    assert not list(tmp_path.glob('scan-*'))


def test_parser_keeps_railpack_runtime_paths_without_forwarding_secrets(tmp_path, monkeypatch):
    import sys

    import jobtology_be.resume.service as module
    runtime_keys = ('MISE_DATA_DIR', 'MISE_CONFIG_DIR', 'MISE_CACHE_DIR', 'MISE_INSTALLS_DIR')
    for key in runtime_keys:
        monkeypatch.setenv(key, '/mise')
    monkeypatch.setenv('JOBTOLOGY_DATABASE_URL', 'synthetic-private-database')
    monkeypatch.setenv('JOBTOLOGY_GOOGLE_CLIENT_SECRET', 'synthetic-private-secret')
    original = asyncio.create_subprocess_exec
    async def native_probe(*args, **kwargs):
        code = '''import json, os, pathlib, sys
assert all(os.environ.get(key) == '/mise' for key in (
    'MISE_DATA_DIR', 'MISE_CONFIG_DIR', 'MISE_CACHE_DIR', 'MISE_INSTALLS_DIR'))
assert 'JOBTOLOGY_DATABASE_URL' not in os.environ
assert 'JOBTOLOGY_GOOGLE_CLIENT_SECRET' not in os.environ
assert pathlib.Path(os.environ['HOME']) == pathlib.Path.cwd()
pathlib.Path(sys.argv[1]).write_text(json.dumps({'ok': True, 'result': {
    'blocks': [], 'warnings': [], 'skipped_pages': []}}))
'''
        return await original(sys.executable, '-c', code, args[-1], **kwargs)
    monkeypatch.setattr(module.asyncio, 'create_subprocess_exec', native_probe)
    service = ResumeService(catalog(), root=tmp_path, formats=('docx',))
    async def stream():
        yield docx_bytes()
    assert asyncio.run(service.scan(USER, 'docx', stream())).blocks == ()
    assert not list(tmp_path.glob('scan-*'))


def test_cancelled_request_erases_partial_upload_and_releases_account_slot(tmp_path):
    service = ResumeService(catalog(), root=tmp_path, formats=('docx',))
    async def stream():
        yield b'private'
        raise asyncio.CancelledError
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(service.scan(USER, 'docx', stream()))
    assert not list(tmp_path.glob('scan-*'))
    with service.slot(USER):
        pass


def test_hwpx_and_pdf_run_through_the_actual_bounded_child(tmp_path):
    import runpy
    from pathlib import Path

    from jobtology_be.document_processor.runtime_check import runtime_report
    probe = runpy.run_path(str(Path(__file__).parents[1] / 'experiments/resume_import/verify_processor.py'))
    service = ResumeService(catalog(), root=tmp_path / 'uploads')
    probe['make_hwpx'](tmp_path / 'sample.hwpx')
    async def scan_file(path):
        async def stream():
            yield path.read_bytes()
        return await service.scan(USER, path.suffix[1:], stream())
    hwpx = asyncio.run(scan_file(tmp_path / 'sample.hwpx'))
    assert 'Python' in hwpx.blocks[0].text
    assert hwpx.blocks[0].page is None
    if not runtime_report()['formats']['pdf']:
        pytest.skip('Native PDF runtime not installed')
    probe['make_pdf'](tmp_path / 'text.pdf')
    probe['make_pdf'](tmp_path / 'scan.pdf', scan=True)
    pdf = asyncio.run(scan_file(tmp_path / 'text.pdf'))
    assert 'Python' in pdf.blocks[0].text and pdf.blocks[0].page == 1
    image_only = asyncio.run(scan_file(tmp_path / 'scan.pdf'))
    assert not image_only.blocks
    assert 'OCR_REQUIRED' in image_only.warnings
    assert image_only.skipped_pages == (1,)
    assert not list(service.root.glob('scan-*'))


def test_disconnect_stops_parser_and_cleans_up_without_waiting_for_timeout(tmp_path, monkeypatch):
    import sys

    import jobtology_be.resume.service as module
    original = asyncio.create_subprocess_exec
    children = []
    async def slow_child(*args, **kwargs):
        process = await original(sys.executable, '-c', 'import time; time.sleep(30)', **kwargs)
        children.append(process)
        return process
    monkeypatch.setattr(module.asyncio, 'create_subprocess_exec', slow_child)
    service = ResumeService(catalog(), root=tmp_path, formats=('docx',))
    async def stream():
        yield docx_bytes()
    async def disconnected():
        return True
    with pytest.raises(ResumeError, match='CLIENT_DISCONNECTED'):
        asyncio.run(service.scan(USER, 'docx', stream(), disconnected=disconnected))
    assert children[0].returncode < 0
    assert not list(tmp_path.glob('scan-*'))


def test_headings_give_category_context_without_creating_fake_qualifications():
    values, _ = literal_suggestions((block('프로젝트'), block(TEXT), block('자격증 · 시험'),
                                    block('TOEIC ' + '9' * 400)), catalog().items())
    assert len(values) == 2
    assert values[0].category == 'project'
    assert values[1].score is None
