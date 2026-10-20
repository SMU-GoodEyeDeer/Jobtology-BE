"""Synthetic DOCX replacements for upstream's local-only example documents."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def docx_without_numbering() -> bytes:
    doc = Document()
    doc.add_paragraph("Synthetic paragraph without a numbering part")
    original = BytesIO()
    doc.save(original)
    output = BytesIO()
    with ZipFile(original) as source, ZipFile(output, "w") as target:
        for item in source.infolist():
            if item.filename == "word/numbering.xml":
                continue
            data = source.read(item.filename)
            if item.filename in {"word/_rels/document.xml.rels", "[Content_Types].xml"}:
                root = ET.fromstring(data)
                for child in list(root):
                    if (child.get("Type", "").endswith("/numbering")
                            or child.get("PartName") == "/word/numbering.xml"):
                        root.remove(child)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    return output.getvalue()


def docx_with_symbol_bullets_and_merges() -> bytes:
    doc = Document()
    numbering = doc.part.numbering_part.element
    for level in numbering.iter(qn("w:lvl")):
        fmt = level.find(qn("w:numFmt"))
        if fmt is None or fmt.get(qn("w:val")) != "bullet":
            continue
        level.find(qn("w:lvlText")).set(qn("w:val"), "\uf0b7")
        properties = level.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            level.append(properties)
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Symbol")
        fonts.set(qn("w:hAnsi"), "Symbol")
        properties.append(fonts)
    for text in ("Bullet 1", "Bullet 2\nsoft break", "Bullet 3"):
        doc.add_paragraph(text, style="List Bullet")
    for _ in range(2):
        table = doc.add_table(rows=3, cols=2)
        for row, cells in enumerate(table.rows):
            for column, cell in enumerate(cells.cells):
                cell.text = f"Cell {row}:{column}"
        table.cell(1, 1).merge(table.cell(2, 1))
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
