from __future__ import annotations

from hashlib import sha1
from io import BytesIO
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
import zipfile

from jobtology_be.document_processor import (
    DocIR,
    DocumentInput,
    TextEdit,
    apply_document_edits,
    read_document,
)
from jobtology_be.document_processor.core import doc_converter
from jobtology_be.document_processor.io_utils import infer_doc_type


class LegacyDocSupportTests(unittest.TestCase):
    @staticmethod
    def _build_docx_bytes(*, invalid_list_style: bool = False) -> bytes:
        from docx import Document

        document = Document()
        run = document.add_paragraph().add_run("Converted legacy document")
        run.bold = True
        output = BytesIO()
        document.save(output)
        data = output.getvalue()
        if not invalid_list_style:
            return data

        source = zipfile.ZipFile(BytesIO(data))
        styles_root = ET.fromstring(source.read("word/styles.xml"))
        word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        ET.SubElement(
            styles_root,
            f"{{{word_ns}}}style",
            {
                f"{{{word_ns}}}type": "list",
                f"{{{word_ns}}}styleId": "LegacyList",
            },
        )
        ET.register_namespace("w", word_ns)

        normalized = BytesIO()
        with zipfile.ZipFile(normalized, "w") as target:
            for item in source.infolist():
                payload = (
                    ET.tostring(styles_root, encoding="utf-8", xml_declaration=True)
                    if item.filename == "word/styles.xml"
                    else source.read(item.filename)
                )
                target.writestr(item, payload)
        source.close()
        return normalized.getvalue()

    def test_infers_doc_from_suffix_and_accepts_explicit_bytes_type(self) -> None:
        self.assertEqual(infer_doc_type(Path("legacy.doc"), "auto"), "doc")
        self.assertEqual(infer_doc_type(b"legacy-doc", "doc"), "doc")

        with self.assertRaises(ValueError):
            infer_doc_type(b"legacy-doc", "auto")

    def test_converter_runs_configured_cli_and_normalizes_list_style(self) -> None:
        converted = self._build_docx_bytes(invalid_list_style=True)

        with tempfile.TemporaryDirectory() as tmp_dir:
            source = Path(tmp_dir) / "legacy.doc"
            converter = Path(tmp_dir) / "doc2docx"
            source.write_bytes(b"legacy-doc")
            converter.write_bytes(b"placeholder")

            def fake_run(command, **kwargs):
                Path(command[-1]).write_bytes(converted)
                return subprocess.CompletedProcess(command, 0, "", "")

            with (
                patch.dict(
                    os.environ,
                    {doc_converter.DOC_CONVERTER_ENV_VAR: str(converter)},
                ),
                patch.object(
                    doc_converter.subprocess, "run", side_effect=fake_run
                ) as run,
            ):
                result = doc_converter.convert_doc_to_docx_bytes(source)

        command = run.call_args.args[0]
        self.assertEqual(command[0], str(converter))
        with zipfile.ZipFile(BytesIO(result)) as archive:
            styles_xml = archive.read("word/styles.xml")
        self.assertNotIn(b'type="list"', styles_xml)
        self.assertIn(b'type="numbering"', styles_xml)

    def test_converter_surfaces_cli_failure(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            source = Path(tmp_dir) / "legacy.doc"
            converter = Path(tmp_dir) / "doc2docx"
            source.write_bytes(b"legacy-doc")
            converter.write_bytes(b"placeholder")

            completed = subprocess.CompletedProcess(
                [str(converter)],
                1,
                "",
                "unsupported legacy document",
            )
            with (
                patch.dict(
                    os.environ,
                    {doc_converter.DOC_CONVERTER_ENV_VAR: str(converter)},
                ),
                patch.object(doc_converter.subprocess, "run", return_value=completed),
                self.assertRaisesRegex(
                    doc_converter.DocConversionError,
                    "unsupported legacy document",
                ),
            ):
                doc_converter.convert_doc_to_docx_bytes(source)

    def test_docir_converts_file_object_once_and_preserves_source_type(self) -> None:
        converted = self._build_docx_bytes()

        with patch.object(
            doc_converter,
            "convert_doc_to_docx_bytes",
            return_value=converted,
        ) as convert:
            document = DocIR.from_file(BytesIO(b"legacy-doc"), doc_type="doc")

        self.assertEqual(convert.call_count, 1)
        converted_path = Path(convert.call_args.args[0])
        self.assertEqual(converted_path.suffix, ".doc")
        self.assertEqual(document.source_doc_type, "doc")
        self.assertEqual(document.paragraphs[0].text, "Converted legacy document")
        self.assertTrue(document.paragraphs[0].runs[0].run_style.bold)

    def test_document_input_infers_doc_from_source_name(self) -> None:
        converted = self._build_docx_bytes()

        with patch.object(
            doc_converter,
            "convert_doc_to_docx_bytes",
            return_value=converted,
        ):
            result = read_document(
                document=DocumentInput(
                    source_bytes=b"legacy-doc",
                    source_name="legacy.doc",
                )
            )

        self.assertEqual(result.source_doc_type, "doc")
        self.assertEqual(result.paragraphs[0].text, "Converted legacy document")

    def test_legacy_doc_text_writeback_produces_docx(self) -> None:
        converted = self._build_docx_bytes()

        with tempfile.TemporaryDirectory() as tmp_dir:
            source = Path(tmp_dir) / "legacy.doc"
            requested_output = Path(tmp_dir) / "legacy_edited.doc"
            source.write_bytes(b"legacy-doc")

            with (
                patch.object(
                    doc_converter,
                    "convert_doc_to_docx_bytes",
                    return_value=converted,
                ),
                patch(
                    "jobtology_be.document_processor.edit_engine.convert_doc_to_docx_bytes",
                    return_value=converted,
                ),
            ):
                document = DocIR.from_file(source)
                run = document.paragraphs[0].runs[0]
                result = apply_document_edits(
                    document=DocumentInput(source_path=source),
                    edits=[
                        TextEdit(
                            target_id=run.node_id,
                            expected_text_hash=sha1(
                                run.text.encode("utf-8")
                            ).hexdigest(),
                            new_text="Edited legacy document",
                        )
                    ],
                    output_path=str(requested_output),
                )

            self.assertTrue(result.ok)
            self.assertEqual(Path(result.output_path).suffix, ".docx")
            self.assertTrue(
                any("written back as DOCX" in item for item in result.warnings)
            )
            reparsed = DocIR.from_file(Path(result.output_path), doc_type="docx")
            self.assertEqual(reparsed.paragraphs[0].text, "Edited legacy document")


if __name__ == "__main__":
    unittest.main()
