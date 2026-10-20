from __future__ import annotations

from io import BytesIO
import unittest
import zipfile

from jobtology_be.document_processor import (
    DocIR,
    extract_docx_range,
    extract_hwpx_range,
    extract_range,
    extract_ranges,
)


class RangeExtractorDispatchTests(unittest.TestCase):
    @staticmethod
    def _build_docx_bytes() -> bytes:
        from docx import Document

        docx = Document()
        for text in ("First", "Second", "Third"):
            docx.add_paragraph(text)
        buffer = BytesIO()
        docx.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _build_hwpx_bytes() -> bytes:
        hwpx_bytes = BytesIO()
        with zipfile.ZipFile(hwpx_bytes, "w") as archive:
            archive.writestr(
                "Contents/header.xml",
                '<?xml version="1.0" encoding="UTF-8"?>\n'
                '<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head" '
                'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core" />\n',
            )
            archive.writestr(
                "Contents/section0.xml",
                '<?xml version="1.0" encoding="UTF-8"?>\n'
                '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
                'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">\n'
                "  <hp:p><hp:run><hp:t>First</hp:t></hp:run></hp:p>\n"
                "  <hp:p><hp:run><hp:t>Second</hp:t></hp:run></hp:p>\n"
                "  <hp:p><hp:run><hp:t>Third</hp:t></hp:run></hp:p>\n"
                "</hs:sec>\n",
            )
        return hwpx_bytes.getvalue()

    def test_dispatches_docx_to_docx_extractor(self) -> None:
        source = self._build_docx_bytes()
        paragraphs = DocIR.from_file(source, doc_type="docx").paragraphs
        start, end = paragraphs[0].node_id, paragraphs[1].node_id

        unified = extract_range(source, doc_type="docx", start_node_id=start, end_node_id=end)
        direct = extract_docx_range(source, start_node_id=start, end_node_id=end)

        self.assertEqual(unified, direct)

    def test_dispatches_hwpx_to_hwpx_extractor(self) -> None:
        source = self._build_hwpx_bytes()
        paragraphs = DocIR.from_file(source, doc_type="hwpx").paragraphs
        start, end = paragraphs[0].node_id, paragraphs[1].node_id

        unified = extract_range(source, doc_type="hwpx", start_node_id=start, end_node_id=end)
        direct = extract_hwpx_range(source, start_node_id=start, end_node_id=end)

        self.assertEqual(unified, direct)

    def test_infers_doc_type_from_bytes(self) -> None:
        source = self._build_docx_bytes()
        paragraphs = DocIR.from_file(source, doc_type="docx").paragraphs

        # doc_type="auto" (default) infers DOCX from the ZIP/OOXML content.
        chunk = extract_range(
            source,
            start_node_id=paragraphs[0].node_id,
            end_node_id=paragraphs[1].node_id,
        )
        texts = [p.text for p in DocIR.from_file(chunk, doc_type="docx").paragraphs if p.text.strip()]
        self.assertEqual(texts, ["First", "Second"])

    def test_pdf_not_supported(self) -> None:
        with self.assertRaises(NotImplementedError):
            extract_ranges(b"%PDF-1.4", [("a", "b")], doc_type="pdf")


if __name__ == "__main__":
    unittest.main()
