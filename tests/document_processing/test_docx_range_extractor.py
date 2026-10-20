from __future__ import annotations

from io import BytesIO
import unittest
from unittest.mock import patch
import zipfile

from jobtology_be.document_processor import DocIR, extract_docx_range, extract_docx_ranges
from jobtology_be.document_processor.core.docx_range_extractor import _paragraph_node_id


class DocxRangeExtractorTests(unittest.TestCase):
    PARAGRAPH_TEXTS = ["First", "Second", "Third", "Fourth"]

    @classmethod
    def _build_sample_docx_bytes(cls) -> bytes:
        from docx import Document

        docx = Document()
        for text in cls.PARAGRAPH_TEXTS:
            docx.add_paragraph(text)
        buffer = BytesIO()
        docx.save(buffer)
        return buffer.getvalue()

    @classmethod
    def _build_three_section_docx_bytes(cls) -> bytes:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.shared import Inches

        docx = Document()
        docx.sections[0].left_margin = Inches(1)
        docx.sections[0].right_margin = Inches(1)
        docx.add_paragraph("Section one")

        docx.add_section(WD_SECTION.NEW_PAGE)
        docx.sections[-1].left_margin = Inches(2)
        docx.sections[-1].right_margin = Inches(2)
        docx.add_paragraph("Section two")

        docx.add_section(WD_SECTION.NEW_PAGE)
        docx.sections[-1].left_margin = Inches(3)
        docx.sections[-1].right_margin = Inches(3)
        docx.add_paragraph("Section three")

        buffer = BytesIO()
        docx.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _visible_texts(docx_bytes: bytes) -> list[str]:
        doc = DocIR.from_file(docx_bytes, doc_type="docx")
        return [p.text for p in doc.paragraphs if p.text.strip()]

    def _paragraph(self, docx_bytes: bytes, text: str):
        for paragraph in DocIR.from_file(docx_bytes, doc_type="docx").paragraphs:
            if paragraph.text.strip() == text:
                return paragraph
        raise AssertionError(f"paragraph {text!r} not found")

    @staticmethod
    def _first_section_margins_inches(docx_bytes: bytes) -> tuple[float, float]:
        from docx import Document

        docx = Document(BytesIO(docx_bytes))
        section = docx.sections[0]
        return round(section.left_margin.inches, 2), round(section.right_margin.inches, 2)

    def test_node_ids_match_docir(self) -> None:
        source = self._build_sample_docx_bytes()
        paragraphs = DocIR.from_file(source, doc_type="docx").paragraphs
        # DocIR node ids follow the s1.p{n} block scheme the extractor addresses.
        for index, paragraph in enumerate(paragraphs, start=1):
            self.assertEqual(paragraph.node_id, _paragraph_node_id(index))

    def test_extract_middle_range_keeps_only_that_range(self) -> None:
        source = self._build_sample_docx_bytes()
        extracted = extract_docx_range(
            source,
            start_node_id=self._paragraph(source, "Second").node_id,
            end_node_id=self._paragraph(source, "Third").node_id,
        )
        self.assertEqual(self._visible_texts(extracted), ["Second", "Third"])

    def test_extract_single_paragraph(self) -> None:
        source = self._build_sample_docx_bytes()
        extracted = extract_docx_range(
            source,
            start_node_id=self._paragraph(source, "First").node_id,
            end_node_id=self._paragraph(source, "First").node_id,
        )
        self.assertEqual(self._visible_texts(extracted), ["First"])

    def test_other_parts_preserved(self) -> None:
        source = self._build_sample_docx_bytes()
        extracted = extract_docx_range(
            source,
            start_node_id=self._paragraph(source, "Second").node_id,
            end_node_id=self._paragraph(source, "Third").node_id,
        )
        with zipfile.ZipFile(BytesIO(extracted)) as archive:
            names = archive.namelist()
            self.assertIn("word/styles.xml", names)
            self.assertIn("word/document.xml", names)

    def test_batch_matches_individual(self) -> None:
        source = self._build_sample_docx_bytes()
        ranges = [
            (self._paragraph(source, "First").node_id, self._paragraph(source, "Second").node_id),
            (self._paragraph(source, "Third").node_id, self._paragraph(source, "Fourth").node_id),
        ]
        batch = extract_docx_ranges(source, ranges)
        self.assertEqual(len(batch), 2)
        self.assertEqual(self._visible_texts(batch[0]), ["First", "Second"])
        self.assertEqual(self._visible_texts(batch[1]), ["Third", "Fourth"])

    def test_batch_parses_source_once(self) -> None:
        from docx import Document

        source = self._build_sample_docx_bytes()
        paragraphs = DocIR.from_file(source, doc_type="docx").paragraphs
        ranges = [
            (paragraphs[0].node_id, paragraphs[1].node_id),
            (paragraphs[2].node_id, paragraphs[3].node_id),
        ]

        with patch("docx.Document", wraps=Document) as document_mock:
            batch = extract_docx_ranges(source, ranges)

        self.assertEqual(document_mock.call_count, 1)
        self.assertEqual(len(batch), 2)
        self.assertEqual(self._visible_texts(batch[0]), ["First", "Second"])
        self.assertEqual(self._visible_texts(batch[1]), ["Third", "Fourth"])

    def test_extract_middle_section_preserves_selected_section_layout(self) -> None:
        source = self._build_three_section_docx_bytes()

        extracted = extract_docx_range(
            source,
            start_node_id=self._paragraph(source, "Section two").node_id,
            end_node_id=self._paragraph(source, "Section two").node_id,
        )

        self.assertEqual(self._visible_texts(extracted), ["Section two"])
        self.assertEqual(self._first_section_margins_inches(extracted), (2.0, 2.0))

    @classmethod
    def _build_docx_with_toc_sdt_bytes(cls) -> bytes:
        """Sample doc with a leading ``<w:sdt>`` (a table-of-contents content control)."""
        from docx import Document
        from docx.oxml import OxmlElement

        docx = Document()
        for text in cls.PARAGRAPH_TEXTS:
            docx.add_paragraph(text)

        sdt = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text_el = OxmlElement("w:t")
        text_el.text = "TOC_MARKER"
        run.append(text_el)
        paragraph.append(run)
        content.append(paragraph)
        sdt.append(content)
        docx.element.body.insert(0, sdt)

        buffer = BytesIO()
        docx.save(buffer)
        return buffer.getvalue()

    def test_strips_out_of_range_sdt_toc(self) -> None:
        source = self._build_docx_with_toc_sdt_bytes()
        extracted = extract_docx_range(
            source,
            start_node_id=self._paragraph(source, "Second").node_id,
            end_node_id=self._paragraph(source, "Third").node_id,
        )

        self.assertEqual(self._visible_texts(extracted), ["Second", "Third"])
        with zipfile.ZipFile(BytesIO(extracted)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        # The leading TOC content control (outside the range) must be gone.
        self.assertNotIn("TOC_MARKER", document_xml)
        self.assertNotIn("<w:sdt>", document_xml)

    @staticmethod
    def _minimal_png() -> bytes:
        import struct
        import zlib

        def chunk(kind: bytes, data: bytes) -> bytes:
            body = kind + data
            return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

        ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
        idat = zlib.compress(b"\x00\xff\xff\xff")
        return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")

    @classmethod
    def _build_docx_with_image_bytes(cls) -> bytes:
        from docx import Document
        from docx.shared import Inches

        docx = Document()
        docx.add_paragraph("First")
        picture_paragraph = docx.add_paragraph()
        picture_paragraph.add_run().add_picture(BytesIO(cls._minimal_png()), width=Inches(1))
        docx.add_paragraph("Third")
        buffer = BytesIO()
        docx.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _media_names(docx_bytes: bytes) -> list[str]:
        with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
            return sorted(n for n in archive.namelist() if n.startswith("word/media/"))

    def test_prunes_media_outside_range(self) -> None:
        source = self._build_docx_with_image_bytes()
        self.assertEqual(len(self._media_names(source)), 1)  # sanity: source has the image

        # Range = "First" only, which references no image → media pruned.
        extracted = extract_docx_range(
            source,
            start_node_id=self._paragraph(source, "First").node_id,
            end_node_id=self._paragraph(source, "First").node_id,
        )
        self.assertEqual(self._media_names(extracted), [])

    def test_keeps_media_referenced_by_range(self) -> None:
        source = self._build_docx_with_image_bytes()
        paragraphs = DocIR.from_file(source, doc_type="docx").paragraphs
        # The image lives on the 2nd block; keep the whole document so it survives.
        extracted = extract_docx_range(
            source,
            start_node_id=paragraphs[0].node_id,
            end_node_id=paragraphs[-1].node_id,
        )
        self.assertEqual(len(self._media_names(extracted)), 1)

    def test_unknown_node_id_raises(self) -> None:
        source = self._build_sample_docx_bytes()
        with self.assertRaises(KeyError):
            extract_docx_range(
                source,
                start_node_id=self._paragraph(source, "First").node_id,
                end_node_id="p_does_not_exist",
            )


if __name__ == "__main__":
    unittest.main()
