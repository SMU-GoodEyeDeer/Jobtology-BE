"""Local feasibility probe. Synthetic documents only; no API, DB, or LLM calls.

Run with the backend's locked Python 3.12 environment. The processor source is
part of BE. This is not an upload handler or a production HTML sanitizer.
"""

import argparse
import hashlib
import json
import subprocess
from html import escape
from html.parser import HTMLParser
from importlib.metadata import version
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)

from jobtology_be.document_processor import (
    DocIR,
    DocumentInput,
    ParagraphIR,
    RunIR,
    RunStyleInfo,
    TextAnnotation,
    read_document,
    render_review_html,
)
from jobtology_be.document_processor.core.doc_converter import _converter_command
from jobtology_be.document_processor.core.hwp_converter import _ensure_jvm

EVIDENCE = "Python과 FastAPI로 로그인 API를 직접 구현하고 PostgreSQL을 연동했습니다."
TABLE_EVIDENCE = "프로젝트에서 REST API 테스트를 작성하고 인증 오류를 수정했습니다."
PDF_EVIDENCE = "I implemented a login API using Python and FastAPI."


def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("이력서 검토 예시 · 합성 데이터", 0)
    doc.add_paragraph("실제 지원자 정보가 없는 로컬 기능 검증 문서입니다.")
    doc.add_heading("프로젝트", 1)
    paragraph = doc.add_paragraph()
    # A quote crossing run boundaries still needs one exact paragraph anchor.
    paragraph.add_run("Python과 Fast")
    paragraph.add_run("API").bold = True
    paragraph.add_run("로 로그인 API를 직접 구현하고 PostgreSQL을 연동했습니다.")
    doc.add_paragraph("기간: 2026.03 ~ 2026.06 · 역할: 인증 API 구현")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "수행 근거"
    table.cell(0, 1).text = TABLE_EVIDENCE
    doc.add_heading("자격증 · 시험", 1)
    doc.add_paragraph("정보처리기사 · 한국산업인력공단 · 2026.06")
    doc.add_paragraph("TOEIC 850 / 990 · 2026.05")
    doc.add_heading("아직 수행 근거가 아닌 문장", 1)
    doc.add_paragraph("앞으로 Rust와 Kubernetes를 배우고 싶습니다.")
    doc.add_paragraph("팀원은 Docker를 담당했고 저는 로그인 API를 구현했습니다.")
    doc.add_paragraph("JavaScript를 사용했습니다. Java 경험이라고 해석하면 안 됩니다.")
    doc.add_paragraph("Python Python")
    # A parser-readable string need not be visible to the person reviewing it.
    doc.add_paragraph().add_run("HIDDEN_KEYWORD").font.hidden = True
    doc.save(path)


def make_hwpx(path: Path) -> None:
    # Minimal synthetic HWPX container, not a visual fidelity benchmark.
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr(
            "Contents/header.xml",
            '<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head"/>',
        )
        archive.writestr(
            "Contents/section0.xml",
            '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
            'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
            f'<hp:p><hp:run><hp:t>{xml_escape(EVIDENCE)}</hp:t></hp:run></hp:p></hs:sec>',
        )


def make_pdf(path: Path, *, scan: bool = False) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    if scan:
        image = DecodedStreamObject()
        image.set_data(b"\xff" * (100 * 100))
        image.update({
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(100),
            NameObject("/Height"): NumberObject(100),
            NameObject("/ColorSpace"): NameObject("/DeviceGray"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/XObject"): DictionaryObject({NameObject("/Im1"): writer._add_object(image)}),
        })
        stream.set_data(b"q 612 0 0 792 0 0 cm /Im1 Do Q")
    else:
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        content = (
            "BT /F1 12 Tf 72 740 Td (Synthetic resume - no personal data) Tj "
            f"0 -28 Td ({PDF_EVIDENCE}) Tj ET"
        )
        stream.set_data(content.encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def contexts(doc: DocIR):
    """Page through read_document, including text nested inside table cells."""
    result = []
    start = 0
    while True:
        page = read_document(document=DocumentInput(doc_ir=doc), start=start, limit=50)
        result.extend(page.paragraphs)
        if page.next_start is None:
            return result
        start = page.next_start


class AttributeProbe(HTMLParser):
    def __init__(self):
        super().__init__()
        self.event_attributes = []

    def handle_starttag(self, tag, attrs):
        self.event_attributes.extend(key for key, _ in attrs if key.lower().startswith("on"))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--native-smoke", action="store_true", help="Also load HWP classes and boot the legacy DOC CLI.")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    source = args.out / "synthetic-resume.docx"
    make_docx(source)
    doc = DocIR.from_file(source)
    paragraphs = contexts(doc)
    by_text = {paragraph.text: paragraph for paragraph in paragraphs}
    assert EVIDENCE in by_text and TABLE_EVIDENCE in by_text
    assert [p.node_id for p in paragraphs] == [p.node_id for p in contexts(DocIR.from_file(source))]
    annotations = [TextAnnotation(
        target_kind="paragraph", target_id=by_text[quote].node_id,
        selected_text=quote, label="수행 근거 · 사용자 확인 필요", color="#D1FAE5",
        note="문장의 존재만 확인했습니다. 역량·수행 수준은 자동 저장되지 않습니다.",
    ) for quote in (EVIDENCE, TABLE_EVIDENCE)]
    review = render_review_html(document=DocumentInput(doc_ir=doc), annotations=annotations)
    assert review.ok and review.html and len(review.resolved_annotations) == 2
    assert "<mark " in review.html
    (args.out / "review.html").write_text(
        review.html.replace('<meta charset="utf-8">',
                            '<meta charset="utf-8"><meta name="darkreader-lock">'),
        encoding="utf-8",
    )

    ambiguous = TextAnnotation(target_kind="paragraph", target_id=by_text["Python Python"].node_id,
                               selected_text="Python", label="keyword")
    ambiguity = render_review_html(document=DocumentInput(doc_ir=doc), annotations=[ambiguous])
    assert not ambiguity.ok
    assert ambiguity.validation.issues[0].code == "selected_text_ambiguous"
    resolved = render_review_html(document=DocumentInput(doc_ir=doc),
                                 annotations=[ambiguous.model_copy(update={"occurrence_index": 1})])
    assert resolved.ok and resolved.resolved_annotations[0].start == 7
    invented = render_review_html(document=DocumentInput(doc_ir=doc), annotations=[
        ambiguous.model_copy(update={"selected_text": "I operated production Kubernetes"}),
    ])
    assert not invented.ok and invented.validation.issues[0].code == "selected_text_not_found"

    hwpx = args.out / "synthetic-resume.hwpx"
    make_hwpx(hwpx)
    hwpx_doc = DocIR.from_file(hwpx)
    hwpx_paragraph = next(p for p in contexts(hwpx_doc) if p.text == EVIDENCE)
    hwpx_review = render_review_html(document=DocumentInput(doc_ir=hwpx_doc), annotations=[
        TextAnnotation(target_kind="paragraph", target_id=hwpx_paragraph.node_id,
                       selected_text=EVIDENCE, label="source evidence"),
    ])
    assert hwpx_review.ok

    pdf = args.out / "synthetic-resume.pdf"
    make_pdf(pdf)
    pdf_doc = DocIR.from_file(pdf, config={"image_output": "off"})
    pdf_paragraph = next(p for p in contexts(pdf_doc) if PDF_EVIDENCE in p.text)
    pdf_review = render_review_html(document=DocumentInput(doc_ir=pdf_doc), annotations=[
        TextAnnotation(target_kind="paragraph", target_id=pdf_paragraph.node_id,
                       selected_text=PDF_EVIDENCE, label="source evidence"),
    ])
    assert pdf_review.ok and pdf_paragraph.page_number == 1

    scan = args.out / "synthetic-scan.pdf"
    make_pdf(scan, scan=True)
    scan_doc = DocIR.from_file(scan, config={"image_output": "off"})
    assert not any(p.text.strip() for p in contexts(scan_doc))
    assert scan_doc.pages[0].parse_skip_reason == "pdf_scan_like"

    # Test the trust boundary at the IR renderer; never execute this HTML.
    unsafe = DocIR(doc_id="synthetic-style-probe", paragraphs=[ParagraphIR(content=[RunIR(
        text="safe text", run_style=RunStyleInfo(color='red" onmouseover="probe'),
    )])])
    probe = AttributeProbe()
    probe.feed(unsafe.to_html())
    styled_review = render_review_html(document=DocumentInput(doc_ir=unsafe), annotations=[])
    review_probe = AttributeProbe()
    review_probe.feed(styled_review.html or "")

    # Check a native font name separately; do not assume the IR probe is reachable.
    native_style_path = args.out / "synthetic-native-style.docx"
    native_style = Document()
    native_style.add_paragraph().add_run("native style probe").font.name = 'Arial" onmouseover="probe'
    native_style.save(native_style_path)
    native_probe = AttributeProbe()
    native_probe.feed(DocIR.from_file(native_style_path).to_html())

    native_smoke = None
    if args.native_smoke:
        import jpype

        _ensure_jvm()
        for name in ("kr.dogfoot.hwplib.reader.HWPReader", "kr.dogfoot.hwp2hwpx.Hwp2Hwpx",
                     "kr.dogfoot.hwpxlib.writer.HWPXWriter"):
            jpype.JClass(name)
        boot = subprocess.run(_converter_command(), capture_output=True, text=True,
                              timeout=10, check=False)
        assert boot.returncode == 64 and "Usage: doc2docx" in boot.stderr + boot.stdout
        native_smoke = {"hwp_converter_classes_loaded": True, "doc_cli_booted": True}

    report = {
        "synthetic_only": True, "llm_calls": 0, "database_writes": 0,
        "processor_upstream_revision": "7541e87",
        "native_smoke": native_smoke,
        "dependencies": {name: version(name) for name in (
            "jobtology-be", "pydantic", "python-docx", "pypdf", "pypdfium2", "jpype1",
        )},
        "docx_nested_table_and_cross_run_quotes": True,
        "docx_reparse_node_ids_stable": True,
        "exact_annotation_quotes": [item.model_dump() for item in review.resolved_annotations],
        "repeated_quote_requires_occurrence": True, "invented_quote_rejected": True,
        "hwpx_text_and_annotation": True, "text_pdf_page_and_annotation": True,
        "scanned_pdf_skipped_not_ocr": True,
        "hidden_text_in_read_document": "HIDDEN_KEYWORD" in by_text,
        "hidden_text_in_regular_html": "HIDDEN_KEYWORD" in doc.to_html(),
        "hidden_text_in_review_html": "HIDDEN_KEYWORD" in review.html,
        "raw_ir_style_event_attributes": probe.event_attributes,
        "review_ir_style_event_attributes": review_probe.event_attributes,
        "native_docx_font_event_attributes": native_probe.event_attributes,
        "document_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "source_spans": [
            {"node_id": p.node_id, "text_hash": p.text_hash,
             "page_number": p.page_number, "text": p.text}
            for p in paragraphs if p.text.strip()
        ],
        "note": "IR style probe is not proof that every native format accepts that payload. "
                "The renderer is not a sanitization boundary. No proficiency was inferred.",
    }
    (args.out / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2),
                                         encoding="utf-8")
    # A static navigation page for the synthetic-only artifacts.
    (args.out / "index.html").write_text(
        '<!doctype html><html lang="ko"><meta charset="utf-8">'
        '<title>Resume import feasibility</title><body>'
        '<h1>이력서 가져오기 · 로컬 검증</h1>'
        '<p>합성 문서만 사용했습니다. AI 호출·DB 저장은 없습니다.</p>'
        '<a href="review.html">원문 근거 하이라이트 보기</a>'
        f'<pre>{escape(json.dumps({k: v for k, v in report.items() if k != "source_spans"}, ensure_ascii=False, indent=2))}</pre>'
        '</body></html>', encoding="utf-8",
    )
    print(json.dumps({key: value for key, value in report.items()
                      if key not in {"source_spans", "exact_annotation_quotes"}}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
